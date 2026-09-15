from io import BytesIO
from django.core.files.base import ContentFile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
from .models import StudentDocument

def _sections(r):
    return {"experiences":r.experiences.all(),"education":r.education_items.all(),"awards":r.awards.all(),"projects":r.projects.all()}

def build_docx(user,r):
    d=Document(); sec=d.sections[0]; compact=r.template=="compact"; sec.top_margin=sec.bottom_margin=Inches(.4 if compact else .55); sec.left_margin=sec.right_margin=Inches(.6 if compact else .7)
    title=d.add_paragraph(); title.alignment=1; run=title.add_run(user.get_full_name() or user.username); run.bold=True; run.font.size=Pt(20)
    contact=" | ".join(x for x in [r.city_state,r.phone,user.personal_email or user.email,r.website] if x); p=d.add_paragraph(contact); p.alignment=1
    def heading(text):
        p=d.add_paragraph(); p.paragraph_format.space_before=Pt(7); p.paragraph_format.space_after=Pt(2); run=p.add_run(text.upper() if r.template=="compact" else text); run.bold=True; run.font.size=Pt(11 if compact else 13);
        if r.template=="modern": run.font.color.rgb=RGBColor(31,78,121)
    if r.objective: heading("Objective"); d.add_paragraph(r.objective)
    if r.skill_list: heading("Skills"); d.add_paragraph(" • ".join(r.skill_list))
    if r.experiences.exists():
        heading("Experience")
        for x in r.experiences.all():
            p=d.add_paragraph(); a=p.add_run(f"{x.title} — {x.company}"); a.bold=True; details=" | ".join(v for v in [x.city_state, " - ".join(v for v in [x.date_from,x.date_to] if v)] if v);
            if details: p.add_run(f"\n{details}");
            if x.summary:d.add_paragraph(x.summary,style="List Bullet")
    if r.education_items.exists():
        heading("Education")
        for x in r.education_items.all():
            p=d.add_paragraph(); p.add_run(x.school_name).bold=True; degree=", ".join(v for v in [x.degree,x.major] if v); dates=" - ".join(v for v in [x.year_from,x.year_to] if v); details=" | ".join(v for v in [degree,dates] if v);
            if details: p.add_run(f"\n{details}");
            if x.description:d.add_paragraph(x.description)
    if r.projects.exists():
        heading("Projects")
        for x in r.projects.all(): d.add_paragraph(f"{x.name}: {x.description}" + (f" ({x.technologies})" if x.technologies else ""))
    if r.awards.exists():
        heading("Awards & Certificates")
        for x in r.awards.all(): d.add_paragraph(" — ".join(v for v in [x.title,x.issuer,x.date] if v) + (f"\n{x.description}" if x.description else ""))
    out=BytesIO(); d.save(out); return out.getvalue()

def build_pdf(user,r):
    out=BytesIO(); styles=getSampleStyleSheet(); accent=colors.HexColor("#1F4E79") if r.template=="modern" else colors.black; compact=r.template=="compact"; styles.add(ParagraphStyle(name="Name",parent=styles["Title"],alignment=TA_CENTER,fontSize=17 if compact else 19,spaceAfter=3)); styles.add(ParagraphStyle(name="Section",parent=styles["Heading2"],fontSize=10 if compact else 11,textColor=accent,spaceBefore=5 if compact else 8,spaceAfter=3))
    story=[Paragraph(user.get_full_name() or user.username,styles["Name"]),Paragraph(" | ".join(x for x in [r.city_state,r.phone,user.personal_email or user.email,r.website] if x),styles["Normal"]),Spacer(1,6)]
    def section(title): story.extend([Paragraph(title,styles["Section"]),HRFlowable(width="100%",thickness=.5),Spacer(1,3)])
    if r.objective: section("OBJECTIVE"); story.append(Paragraph(r.objective,styles["BodyText"]))
    if r.skill_list: section("SKILLS"); story.append(Paragraph(" • ".join(r.skill_list),styles["BodyText"]))
    if r.experiences.exists():
        section("EXPERIENCE")
        for x in r.experiences.all(): story.extend([Paragraph(f"<b>{x.title}</b> — {x.company}",styles["BodyText"]),Paragraph(f"{x.city_state} | {x.date_from} - {x.date_to}",styles["Normal"]),Paragraph(x.summary,styles["BodyText"]),Spacer(1,4)])
    if r.education_items.exists():
        section("EDUCATION")
        for x in r.education_items.all(): story.extend([Paragraph(f"<b>{x.school_name}</b> — {x.degree} {x.major}",styles["BodyText"]),Paragraph(f"{x.city_state} | {x.year_from} - {x.year_to}",styles["Normal"]),Paragraph(x.description,styles["BodyText"]),Spacer(1,4)])
    if r.projects.exists():
        section("PROJECTS")
        for x in r.projects.all(): story.append(Paragraph(f"<b>{x.name}</b>: {x.description} {x.technologies}",styles["BodyText"]))
    if r.awards.exists():
        section("AWARDS & CERTIFICATES")
        for x in r.awards.all(): story.append(Paragraph(" — ".join(v for v in [x.title,x.issuer,x.date] if v)+(f": {x.description}" if x.description else ""),styles["BodyText"]))
    SimpleDocTemplate(out,pagesize=LETTER,rightMargin=38 if compact else 45,leftMargin=38 if compact else 45,topMargin=28 if compact else 36,bottomMargin=28 if compact else 36).build(story); return out.getvalue()

def save_to_hub(user,r,fmt,data):
    doc=StudentDocument(student=user,title=f"{r.name} ({fmt.upper()})",kind=StudentDocument.Kind.RESUME)
    doc.file.save(f"resume_{r.pk}_{r.name.replace(' ','_')}.{fmt}",ContentFile(data),save=True); return doc
